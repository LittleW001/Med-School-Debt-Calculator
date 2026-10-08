#!/usr/bin/env python3
"""Estimate medical-school debt, training-year repayment, and attending capacity.

Input workbook columns, in order:
    School Name | Semester Tuition + Fees | Annual Cost of Living |
    Annual Financial Aid / Scholarships

Rates are decimals: use 0.055 for 5.5%.
Good/Median/Bad are editable planning cases, not statistical quartiles.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

try:
    import matplotlib.pyplot as plt
    import pandas as pd
    from matplotlib.ticker import FuncFormatter
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
except ImportError as exc:
    raise SystemExit(
        "Install the required packages with: "
        "python -m pip install pandas openpyxl matplotlib"
    ) from exc


ROOT = Path(__file__).parent
INPUT_FILE = ROOT / "input" / "medical_school_costs_template.xlsx"
OUTPUT_DIR = ROOT / "output"


# EDIT THESE SETTINGS ---------------------------------------------------------

PROGRAM_YEARS = 4
SEMESTERS_PER_YEAR = 2
PAYOFF_TERMS_MONTHS = (24, 96, 192)
# Longest payoff duration the ideal-term search will consider.
MAX_IDEAL_TERM_MONTHS = 360

# Tax rates are planning assumptions, not tax-return calculations.
RESIDENT_EFFECTIVE_TAX_RATE = 0.28
ATTENDING_EFFECTIVE_TAX_RATE = 0.35


def living_expenses() -> float:
    """Return total MONTHLY living expenses. Edit or add variables here."""
    housing = 1_000.00
    utilities = 250.00
    food = 400.00
    transportation = 320.00
    insurance = 0
    phone_and_internet = 0.00
    medical_out_of_pocket = 500.00
    other = 250.00
    return sum(
        (
            housing,
            utilities,
            food,
            transportation,
            insurance,
            phone_and_internet,
            medical_out_of_pocket,
            other,
        )
    )


# 2025 AAMC national unweighted mean stipends by PGY level.
RESIDENT_SALARIES_BY_PGY = [68_166, 70_499, 73_301, 77_593, 81_807]

# Median attending total-compensation estimates from SalaryDr's 2026 national
# self-reported data. These are editable and vary materially by location,
# practice setting, workload, and employment model.
CAREERS = {
    "Anesthesiology": {
        "training_years": 4,
        "training_path": "4-year anesthesiology residency",
        "attending_salary": 530_000.00,
    },
    "Endocrinology": {
        "training_years": 5,
        "training_path": "3-year internal medicine residency + 2-year fellowship",
        "attending_salary": 360_000.00,
    },
    "Primary Care": {
        "training_years": 3,
        "training_path": "3-year family medicine residency",
        "attending_salary": 310_000.00,
    },
}

# Federal assumptions generally applicable to new professional-school
# borrowers beginning July 1, 2026. Update if your eligibility differs.
FEDERAL_ANNUAL_LIMIT = 50_000.00
FEDERAL_PROFESSIONAL_AGGREGATE_LIMIT = 200_000.00
FEDERAL_LIFETIME_LIMIT = 257_500.00
FEDERAL_ORIGINATION_FEE = 0.01057
PRIOR_TOTAL_FEDERAL_AMOUNT_RECEIVED = 23_250.00
PRIOR_GRAD_PROF_FEDERAL_AMOUNT_RECEIVED = 0.00
ALLOW_PRIVATE_LOANS_FOR_FUNDING_GAP = True

UNDERGRAD_LOANS = [
    {"name": "Undergrad Loan 1", "principal": 3_500.00, "rate": 0.0550, "accrues": False},
    {"name": "Undergrad Loan 2", "principal": 2_318.47, "rate": 0.0550, "accrues": True},
    {"name": "Undergrad Loan 3", "principal": 7_315.78, "rate": 0.0653, "accrues": True},
    {"name": "Undergrad Loan 4", "principal": 7_942.18, "rate": 0.0639, "accrues": True},
    {"name": "Undergrad Loan 5", "principal": 3_774.76, "rate": 0.0652, "accrues": True},
    {"name": "Undergrad Loan 6", "principal": 3_774.76, "rate": 0.0700, "accrues": True},
]

MED_SCHOOL_SCENARIOS = {
    "Good": {
        "tuition_growth": 0.02,
        "living_growth": 0.02,
        "federal_rates": [0.0650] * 4,
        "private_rate": 0.0750,
        "private_fee": 0.0,
    },
    "Median": {
        "tuition_growth": 0.04,
        "living_growth": 0.03,
        "federal_rates": [0.0807] * 4,
        "private_rate": 0.1100,
        "private_fee": 0.0,
    },
    "Bad": {
        "tuition_growth": 0.06,
        "living_growth": 0.05,
        "federal_rates": [0.0950] * 4,
        "private_rate": 0.1500,
        "private_fee": 0.0,
    },
}


# CALCULATIONS ----------------------------------------------------------------

def monthly_payment(balance: float, annual_rate: float, months: int) -> float:
    if balance <= 0:
        return 0.0
    monthly_rate = annual_rate / 12
    if monthly_rate == 0:
        return balance / months
    return balance * monthly_rate / (1 - (1 + monthly_rate) ** -months)


def payment_profile(
    components: list[tuple[float, float]], months: int
) -> tuple[float, float, float]:
    payment = sum(monthly_payment(balance, rate, months) for balance, rate in components)
    principal = sum(balance for balance, _ in components)
    total_paid = payment * months
    return payment, max(0.0, total_paid - principal), total_paid


def ideal_profile(
    components: list[tuple[float, float]], monthly_budget: float
) -> tuple[int | None, float | None, float | None]:
    if monthly_budget <= 0:
        return None, None, None
    for months in range(12, MAX_IDEAL_TERM_MONTHS + 1):
        payment, interest, _ = payment_profile(components, months)
        if payment <= monthly_budget:
            return months, payment, interest
    return None, None, None


def monthly_loan_capacity(annual_salary: float, tax_rate: float) -> float:
    """Upper-bound payment after estimated taxes and listed living expenses."""
    return max(0.0, annual_salary * (1 - tax_rate) / 12 - living_expenses())


def gross_salary_needed(monthly_loan_payment: float, tax_rate: float) -> float:
    """Gross salary needed to cover living expenses plus a monthly loan payment."""
    return 12 * (living_expenses() + monthly_loan_payment) / (1 - tax_rate)


def interest_only_payment(components: list[tuple[float, float]]) -> float:
    """Monthly amount needed to prevent balance growth under this model."""
    return sum(balance * rate / 12 for balance, rate in components)


def payoff_months(
    components: list[tuple[float, float]],
    monthly_payment: float,
    max_months: int = 1_200,
) -> int | None:
    """Return months to payoff using a fixed payment and highest-rate-first order."""
    balances = [[balance, rate] for balance, rate in components if balance > 0.005]
    if not balances:
        return 0
    if monthly_payment <= interest_only_payment(components):
        return None

    for month in range(1, max_months + 1):
        for loan in balances:
            loan[0] += loan[0] * loan[1] / 12

        payment_left = min(monthly_payment, sum(balance for balance, _ in balances))
        for loan in sorted(balances, key=lambda item: item[1], reverse=True):
            applied = min(payment_left, loan[0])
            loan[0] -= applied
            payment_left -= applied
            if payment_left <= 0:
                break

        if sum(balance for balance, _ in balances) <= 0.005:
            return month
    return None


def resident_salary(pgy: int) -> float:
    if pgy <= len(RESIDENT_SALARIES_BY_PGY):
        return RESIDENT_SALARIES_BY_PGY[pgy - 1]
    return RESIDENT_SALARIES_BY_PGY[-1] * 1.03 ** (pgy - len(RESIDENT_SALARIES_BY_PGY))


def simulate_training(
    components: list[tuple[float, float]], years: int
) -> tuple[list[tuple[float, float]], dict]:
    """Accrue monthly interest and apply available payments highest-rate first."""
    balances = [[balance, rate] for balance, rate in components]
    salaries = [resident_salary(pgy) for pgy in range(1, years + 1)]
    capacities = [
        monthly_loan_capacity(salary, RESIDENT_EFFECTIVE_TAX_RATE)
        for salary in salaries
    ]
    total_interest = 0.0
    total_payments = 0.0

    for capacity in capacities:
        for _ in range(12):
            interest = [balance * rate / 12 for balance, rate in balances]
            total_interest += sum(interest)
            for loan, accrued in zip(balances, interest):
                loan[0] += accrued

            payment_left = min(capacity, sum(balance for balance, _ in balances))
            total_payments += payment_left
            for loan in sorted(balances, key=lambda item: item[1], reverse=True):
                applied = min(payment_left, loan[0])
                loan[0] -= applied
                payment_left -= applied
                if payment_left <= 0:
                    break

    ending_components = [(balance, rate) for balance, rate in balances if balance > 0.005]
    return ending_components, {
        "salaries": salaries,
        "capacities": capacities,
        "total_interest": total_interest,
        "total_payments": total_payments,
    }


def career_projection(
    components: list[tuple[float, float]], career: dict
) -> dict:
    start_debt = sum(balance for balance, _ in components)
    start_interest_only = interest_only_payment(components)
    ending_components, training = simulate_training(
        components, career["training_years"]
    )
    end_debt = sum(balance for balance, _ in ending_components)
    end_interest_only = interest_only_payment(ending_components)
    average_capacity = sum(training["capacities"]) / len(training["capacities"])
    attending_capacity = monthly_loan_capacity(
        career["attending_salary"], ATTENDING_EFFECTIVE_TAX_RATE
    )
    ideal_months, ideal_payment, ideal_interest = ideal_profile(
        ending_components, attending_capacity
    )

    def coverage(payment: float, minimum: float) -> float | None:
        return payment / minimum if minimum > 0 else None

    result = {
        "Ideal Payment Budget": attending_capacity,
        "Ideal Duration (Months)": ideal_months,
        "Ideal Monthly Payment": ideal_payment,
        "Ideal Repayment Interest": ideal_interest,
        "Minimum Salary Needed for Ideal Payment": (
            gross_salary_needed(ideal_payment, ATTENDING_EFFECTIVE_TAX_RATE)
            if ideal_payment
            else None
        ),
        "Training Path": career["training_path"],
        "Resident/Fellow Years": career["training_years"],
        "PGY-1 Gross Salary": training["salaries"][0],
        "Average Annual Training Salary": sum(training["salaries"]) / len(training["salaries"]),
        "Total Training Gross Income": sum(training["salaries"]),
        "Monthly Living Expenses": living_expenses(),
        "Average Monthly Resident Loan Capacity": average_capacity,
        "Starting Interest-Only Payment": start_interest_only,
        "Resident Capacity (% of Interest-Only)": coverage(
            average_capacity, start_interest_only
        ),
        "Total Resident/Fellow Loan Payments": training["total_payments"],
        "Interest Accrued During Training": training["total_interest"],
        "Debt After Training": end_debt,
        "Debt Change During Training": end_debt - start_debt,
        "Post-Training Interest-Only Payment": end_interest_only,
        "Median Attending Salary Assumption": career["attending_salary"],
        "Attending Monthly Loan Capacity": attending_capacity,
        "Attending Capacity (% of Interest-Only)": coverage(
            attending_capacity, end_interest_only
        ),
        "Salary Needed: Interest Only": gross_salary_needed(
            end_interest_only, ATTENDING_EFFECTIVE_TAX_RATE
        ),
    }
    for percent in (0.05, 0.10, 0.20):
        target_payment = end_interest_only + end_debt * percent / 12
        result[f"Salary Needed: Interest + {percent:.0%} Balance/Year"] = (
            gross_salary_needed(target_payment, ATTENDING_EFFECTIVE_TAX_RATE)
        )
        result[f"Months to Debt-Free: Interest + {percent:.0%} Balance/Year"] = (
            payoff_months(ending_components, target_payment)
        )
    return result


def validate_settings() -> None:
    rates = [loan["rate"] for loan in UNDERGRAD_LOANS]
    for scenario in MED_SCHOOL_SCENARIOS.values():
        if len(scenario["federal_rates"]) != PROGRAM_YEARS:
            raise ValueError(f"Each scenario needs {PROGRAM_YEARS} federal rates.")
        rates += scenario["federal_rates"] + [
            scenario["tuition_growth"],
            scenario["living_growth"],
            scenario["private_rate"],
            scenario["private_fee"],
        ]
    rates += [RESIDENT_EFFECTIVE_TAX_RATE, ATTENDING_EFFECTIVE_TAX_RATE]
    if any(rate < 0 or rate >= 1 for rate in rates):
        raise ValueError("Enter rates as decimals from 0 up to 1; 5.5% is 0.055.")
    if any(loan["principal"] < 0 for loan in UNDERGRAD_LOANS):
        raise ValueError("Loan principals cannot be negative.")
    if living_expenses() < 0 or MAX_IDEAL_TERM_MONTHS < 12:
        raise ValueError("Living expenses must be nonnegative and the term limit at least 12 months.")


def read_schools(path: Path) -> pd.DataFrame:
    frame = pd.read_excel(path, sheet_name=0)
    if frame.shape[1] < 4:
        raise ValueError("The input workbook needs at least four columns.")
    frame = frame.iloc[:, :4].copy()
    frame.columns = [
        "School Name",
        "Semester Tuition + Fees",
        "Annual Cost of Living",
        "Annual Financial Aid / Scholarships",
    ]
    frame = frame.dropna(subset=["School Name"])
    frame["School Name"] = frame["School Name"].astype(str).str.strip()
    frame = frame[frame["School Name"] != ""]
    for column in frame.columns[1:]:
        frame[column] = pd.to_numeric(frame[column], errors="raise")
    frame["Annual Financial Aid / Scholarships"] = frame[
        "Annual Financial Aid / Scholarships"
    ].fillna(0.0)
    if frame.empty or frame.iloc[:, 1:3].isna().any().any():
        raise ValueError("Tuition and living cost cannot be blank.")
    if (frame.iloc[:, 1:] < 0).any().any():
        raise ValueError("School costs must be nonnegative.")
    return frame


def undergrad_projection() -> tuple[list[tuple[float, float]], float, float]:
    components = []
    principal_now = 0.0
    interest_by_graduation = 0.0
    for loan in UNDERGRAD_LOANS:
        principal = float(loan["principal"])
        rate = float(loan["rate"])
        interest = principal * rate * PROGRAM_YEARS if loan["accrues"] else 0.0
        components.append((principal + interest, rate))
        principal_now += principal
        interest_by_graduation += interest
    return components, principal_now, interest_by_graduation


def estimate_school(
    school: pd.Series, scenario_name: str, scenario: dict
) -> tuple[dict, list[tuple[float, float]]]:
    components, undergrad_principal, undergrad_interest = undergrad_projection()
    professional_room = max(
        0.0,
        FEDERAL_PROFESSIONAL_AGGREGATE_LIMIT
        - PRIOR_GRAD_PROF_FEDERAL_AMOUNT_RECEIVED,
    )
    lifetime_room = max(
        0.0, FEDERAL_LIFETIME_LIMIT - PRIOR_TOTAL_FEDERAL_AMOUNT_RECEIVED
    )
    coa = aid_total = federal_principal = private_principal = 0.0
    med_interest = funding_gap = 0.0

    for year in range(PROGRAM_YEARS):
        annual_federal_room = FEDERAL_ANNUAL_LIMIT
        tuition = school["Semester Tuition + Fees"] * (
            1 + scenario["tuition_growth"]
        ) ** year
        living = school["Annual Cost of Living"] * (
            1 + scenario["living_growth"]
        ) ** year
        aid_per_semester = (
            school["Annual Financial Aid / Scholarships"] / SEMESTERS_PER_YEAR
        )

        for semester in range(SEMESTERS_PER_YEAR):
            semester_cost = tuition + living / SEMESTERS_PER_YEAR
            aid = min(semester_cost, aid_per_semester)
            net_cost = semester_cost - aid
            years_to_graduation = (
                PROGRAM_YEARS - year - semester / SEMESTERS_PER_YEAR
            )
            semesters_left = SEMESTERS_PER_YEAR - semester
            federal_needed = net_cost / (1 - FEDERAL_ORIGINATION_FEE)
            federal = min(
                federal_needed,
                annual_federal_room / semesters_left,
                professional_room,
                lifetime_room,
            )
            federal_net = federal * (1 - FEDERAL_ORIGINATION_FEE)
            remaining_cost = max(0.0, net_cost - federal_net)

            private = 0.0
            if ALLOW_PRIVATE_LOANS_FOR_FUNDING_GAP and remaining_cost:
                private = remaining_cost / (1 - scenario["private_fee"])
            else:
                funding_gap += remaining_cost

            federal_interest = (
                federal * scenario["federal_rates"][year] * years_to_graduation
            )
            private_interest = (
                private * scenario["private_rate"] * years_to_graduation
            )
            if federal:
                components.append(
                    (federal + federal_interest, scenario["federal_rates"][year])
                )
            if private:
                components.append(
                    (private + private_interest, scenario["private_rate"])
                )

            coa += semester_cost
            aid_total += aid
            federal_principal += federal
            private_principal += private
            med_interest += federal_interest + private_interest
            annual_federal_room -= federal
            professional_room -= federal
            lifetime_room -= federal

    debt_at_graduation = sum(balance for balance, _ in components)
    blended_rate = (
        sum(balance * rate for balance, rate in components) / debt_at_graduation
        if debt_at_graduation
        else 0.0
    )
    result = {
        "School Name": school["School Name"],
        "Scenario": scenario_name,
        "4-Year Cost of Attendance": coa,
        "Scholarships Applied": aid_total,
        "Med-School Federal Principal": federal_principal,
        "Med-School Private Principal": private_principal,
        "Funding Gap": funding_gap,
        "Undergrad Principal Now": undergrad_principal,
        "Undergrad In-School Interest": undergrad_interest,
        "Med-School In-School Interest": med_interest,
        "Debt at Graduation": debt_at_graduation,
        "Blended Repayment Rate": blended_rate,
        # Career-specific values overwrite these placeholders after residency.
        "Ideal Payment Budget": None,
        "Ideal Duration (Months)": None,
        "Ideal Monthly Payment": None,
        "Ideal Repayment Interest": None,
        "Minimum Salary Needed for Ideal Payment": None,
    }
    for months in PAYOFF_TERMS_MONTHS:
        payment, interest, total_paid = payment_profile(components, months)
        result[f"{months}-Month Payment"] = payment
        result[f"{months}-Month Repayment Interest"] = interest
        result[f"{months}-Month Total Paid"] = total_paid
    return result, components


# OUTPUT ----------------------------------------------------------------------

def write_workbook(results_by_career: dict[str, pd.DataFrame], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name, results in results_by_career.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(list(results.columns))
        for row in results.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in row])
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 60
        for cell in sheet[1]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions["A"].width = 42
        for column in range(2, sheet.max_column + 1):
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 26
        sheet.column_dimensions["AA"].width = 48
        for column, header in enumerate(results.columns, start=1):
            if "Rate" in header or "(%" in header:
                number_format = "0.0%"
            elif "Years" in header or "Months" in header:
                number_format = "0"
            elif column > 2 and header != "Training Path":
                number_format = '$#,##0.00'
            else:
                continue
            for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
                for cell in cells:
                    cell.number_format = number_format

    workbook.save(output_path)


def make_chart(
    results: pd.DataFrame, output_path: Path, school_name: str | None = None
) -> None:
    school_name = school_name or results.iloc[0]["School Name"]
    selected = results[results["School Name"] == school_name].set_index("Scenario")
    scenarios = [name for name in MED_SCHOOL_SCENARIOS if name in selected.index]
    figure, axes = plt.subplots(1, len(scenarios), figsize=(12, 4.8), sharey=True)
    if len(scenarios) == 1:
        axes = [axes]

    terms = list(PAYOFF_TERMS_MONTHS)
    for axis, scenario_name in zip(axes, scenarios):
        row = selected.loc[scenario_name]
        principal = (
            row["Undergrad Principal Now"]
            + row["Med-School Federal Principal"]
            + row["Med-School Private Principal"]
        )
        school_interest = (
            row["Undergrad In-School Interest"]
            + row["Med-School In-School Interest"]
        )
        repayment_interest = [
            row[f"{months}-Month Repayment Interest"] for months in terms
        ]
        axis.bar(terms, [principal] * 3, width=30, color="0.80", label="Principal")
        axis.bar(
            terms,
            [school_interest] * 3,
            width=30,
            bottom=[principal] * 3,
            color="0.50",
            label="In-school interest",
        )
        axis.bar(
            terms,
            repayment_interest,
            width=30,
            bottom=[principal + school_interest] * 3,
            color="0.15",
            label="Repayment interest",
        )
        axis.set_title(scenario_name)
        axis.set_xlabel("Payoff duration (months)")
        axis.set_xticks(terms)
        axis.grid(axis="y", color="0.90", linewidth=0.8)

    axes[0].set_ylabel("Total dollars paid")
    axes[0].yaxis.set_major_formatter(
        FuncFormatter(lambda value, _: f"${value / 1000:,.0f}k")
    )
    handles, labels = axes[0].get_legend_handles_labels()
    figure.legend(
        handles,
        labels,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.90),
        ncol=3,
        frameon=False,
    )
    figure.suptitle(
        f"Principal and interest by payoff duration\n{school_name}", y=0.99
    )
    figure.tight_layout(rect=(0, 0, 1, 0.80))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def main() -> None:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Could not find {INPUT_FILE}. Put medical_school_costs_template.xlsx "
            "in the input folder beside this script."
        )
    validate_settings()
    schools = read_schools(INPUT_FILE)
    base_rows = []
    career_rows = {name: [] for name in CAREERS}

    for _, school in schools.iterrows():
        for scenario_name, scenario in MED_SCHOOL_SCENARIOS.items():
            base, components = estimate_school(school, scenario_name, scenario)
            base_rows.append(base)
            for career_name, career in CAREERS.items():
                career_rows[career_name].append(
                    {**base, **career_projection(components, career)}
                )

    today = date.today().strftime("%m_%d_%Y")
    excel_output = OUTPUT_DIR / f"med_debt_by_profession_{today}.xlsx"
    chart_output = OUTPUT_DIR / f"debt_payment_comparison_{today}.png"
    frames = {name: pd.DataFrame(rows) for name, rows in career_rows.items()}
    write_workbook(frames, excel_output)
    make_chart(pd.DataFrame(base_rows), chart_output)
    print(f"Created Excel file: {excel_output}")
    print(f"Created chart: {chart_output}")


if __name__ == "__main__":
    main()
